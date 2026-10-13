"""Запись книги XLSX.

Защита от formula injection. Текст из каталога пишется ТОЛЬКО как
строковая ячейка (тип `s`), даже если он похож на формулу (`=...`) или
код ошибки (`#N/A`): openpyxl иначе превратил бы такую строку в формулу.
Ячейкам, которые начинаются с `= + - @` (и табуляции/возврата каретки),
дополнительно ставится признак quotePrefix — Excel показывает их как
текст и не пересчитывает даже после редактирования ячейки.

Значение при этом не меняется (апостроф в данные не добавляется),
поэтому выгруженный файл возвращается импортом без искажений.

Символы, недопустимые в XML (управляющие), в XLSX записать нельзя в
принципе — они удаляются. Импорт такие символы и не принимает.
"""

from __future__ import annotations

from collections.abc import Sequence
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE, Cell, MergedCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from app.modules.imports.exceptions import ExportValueTooLongError
from app.modules.imports.xlsx_format import EXCEL_CELL_MAX_LENGTH, README_SHEET, SheetSpec

FORMULA_TRIGGERS = ("=", "+", "-", "@", "\t", "\r")

_HEADER_FONT = Font(bold=True)
_HEADER_FILL = PatternFill("solid", start_color="FFE8EDF3", end_color="FFE8EDF3")


def is_formula_like(text: str) -> bool:
    return text.startswith(FORMULA_TRIGGERS)


def write_cell(cell: Cell | MergedCell, value: Any) -> None:
    """Типобезопасная запись значения в ячейку."""
    if isinstance(cell, MergedCell):
        raise TypeError("Запись в объединённую ячейку")
    if value is None:
        return
    if isinstance(value, str):
        text = ILLEGAL_CHARACTERS_RE.sub("", value)
        if len(text) > EXCEL_CELL_MAX_LENGTH:
            raise ExportValueTooLongError(
                details={"sheet": cell.parent.title, "cell": cell.coordinate}
            )
        cell.value = text
        # Строго строка: не формула и не код ошибки Excel.
        cell.data_type = "s"
        if is_formula_like(text):
            cell.quotePrefix = True
        return
    if isinstance(value, bool | int | Decimal):
        cell.value = value
        return
    raise TypeError(f"Неподдерживаемый тип значения для XLSX: {type(value).__name__}")


def _write_sheet(
    worksheet: Worksheet,
    spec: SheetSpec,
    rows: Sequence[Sequence[Any]],
    choices: dict[str, list[str]],
) -> None:
    for index, column in enumerate(spec.columns, start=1):
        cell = worksheet.cell(row=1, column=index, value=column.name)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        worksheet.column_dimensions[get_column_letter(index)].width = column.width
    worksheet.freeze_panes = "A2"

    for row_index, row in enumerate(rows, start=2):
        for column_index, value in enumerate(row, start=1):
            write_cell(worksheet.cell(row=row_index, column=column_index), value)

    # Выпадающие списки для перечислений — подсказка при ручном вводе.
    # Импорт всё равно проверяет значения сам.
    for name, values in choices.items():
        if name not in spec.column_names:
            continue
        letter = get_column_letter(spec.column_names.index(name) + 1)
        validation = DataValidation(type="list", formula1=f'"{",".join(values)}"', allow_blank=True)
        validation.add(f"{letter}2:{letter}1048576")
        worksheet.add_data_validation(validation)


def write_workbook(
    sheets: Sequence[tuple[SheetSpec, Sequence[Sequence[Any]]]],
    *,
    readme: Sequence[str],
    choices: dict[str, dict[str, list[str]]] | None = None,
) -> bytes:
    workbook = Workbook()
    default = workbook.active
    if default is not None:
        workbook.remove(default)

    for spec, rows in sheets:
        worksheet = workbook.create_sheet(spec.name)
        _write_sheet(worksheet, spec, rows, (choices or {}).get(spec.name, {}))

    readme_sheet = workbook.create_sheet(README_SHEET)
    readme_sheet.column_dimensions["A"].width = 120
    for index, line in enumerate(readme, start=1):
        cell = readme_sheet.cell(row=index, column=1)
        write_cell(cell, line)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    if readme:
        readme_sheet.cell(row=1, column=1).font = _HEADER_FONT

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
