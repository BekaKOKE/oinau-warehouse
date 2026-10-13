"""Шаблон книги, пример архива и отчёт пакетного импорта.

Шаблон и пример собираются на лету: в репозитории нет бинарных файлов,
которые разъехались бы с кодом разбора. Пример архива проходит проверку
архива и файлов как есть; категорию в книге замените на свою — иначе
предпросмотр покажет, что такой категории нет.

Все значения пишутся через `write_cell` (xlsx_writer.py): защита от
formula injection та же, что у экспорта каталога.
"""

from __future__ import annotations

import zipfile
from collections.abc import Sequence
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw

from app.modules.imports.bundle_format import BUNDLE_XLSX_NAME, PHOTOS_DIR, SCHEMES_DIR
from app.modules.imports.xlsx_writer import write_cell

_HEADER_FONT = Font(bold=True)
_HEADER_FILL = PatternFill("solid", start_color="FFE8EDF3", end_color="FFE8EDF3")

# Колонки шаблона в порядке следования: (заголовок, ширина).
TEMPLATE_COLUMNS: tuple[tuple[str, int], ...] = (
    ("Артикул", 16),
    ("Название", 40),
    ("Категория", 34),
    ("Описание", 50),
    ("Цена от 1", 12),
    ("Цена от 100", 12),
    ("Цена от 500", 12),
    ("Цена от 1000", 12),
    ("Валюта", 8),
    ("Остаток", 10),
    ("Фото", 40),
    ("Схема", 24),
    ("хар: Цвет", 14),
    ("хар: Материал", 18),
    ("Код 1С", 16),
    ("Активен", 9),
    ("Избранное", 10),
)

README_LINES: tuple[str, ...] = (
    "Массовый импорт товаров OINAU PRO: ZIP-архив с книгой и файлами",
    "",
    "Структура архива:",
    f"  {BUNDLE_XLSX_NAME} — эта книга (лист «Товары»);",
    f"  {PHOTOS_DIR}/ — фотографии товаров (jpg, jpeg, png, webp, до 15 МБ каждая);",
    f"  {SCHEMES_DIR}/ — схемы и чертежи (pdf, dwg, dxf, jpg, png, до 50 МБ каждый).",
    "",
    "Правила:",
    "• Ключ товара — артикул. Товар с существующим артикулом обновляется, новый — создаётся.",
    "• Обязательны только «Артикул» и, для нового товара, «Название».",
    "• Пустая ячейка или отсутствующая колонка — значение не меняется "
    "(у нового товара — значение по умолчанию: опубликован, активен, шт.).",
    "• Категория — название («Комплектующие для качелей»), путь «Качели / Сиденья» "
    "или slug. Несколько категорий — через «;», первая — основная.",
    "• Цена: «Цена» (одна цена) или «Цена от N» — цены по количеству. "
    "Колонки «Цена от …» можно добавлять и удалять.",
    "• Валюта — только валюта прайс-листа (KZT); колонку можно не заполнять.",
    "• Остаток — количество на основном складе.",
    "• Фото: имена файлов из папки photos через «;» (504-176.jpg;504-176-2.jpg) "
    "или колонки «Фото 1», «Фото 2», … Первое фото — обложка.",
    "• Заполненная ячейка «Фото» задаёт полный набор фото товара; фото, которых нет "
    "в списке, снимаются с товара (остаются в медиатеке). Уже загруженные фото "
    "можно не класть в архив повторно — достаточно указать их имена.",
    "• Схема: имя файла из папки schemes (504-176.pdf). Схема показывается на "
    "странице товара отдельным блоком, не в галерее.",
    "• Характеристики: колонки «хар: Название» (хар: Цвет) или одна колонка "
    "«Характеристики» в виде «Цвет: чёрный; Материал: резина».",
    "• Колонки-пометки начинайте с «#» — импорт их пропускает.",
    "• Формулы не принимаются: вставьте значения («Специальная вставка» -> «Значения»).",
    "",
    "Порядок: «Загрузить ZIP» -> проверка -> предпросмотр изменений -> «Применить». "
    "До подтверждения каталог не меняется.",
)

EXAMPLE_ROWS: tuple[tuple[Any, ...], ...] = (
    (
        "504-176",
        "Круглое сиденье без подвесов",
        "Комплектующие для качелей",
        "Сиденье из армированной резины для качелей.",
        3900,
        3650,
        3400,
        3250,
        "KZT",
        8376,
        "504-176.jpg;504-176-2.jpg",
        "504-176.pdf",
        "чёрный",
        "резина",
        None,
        "да",
        "нет",
    ),
    (
        "1639000R",
        "Сиденье",
        "Комплектующие для качелей",
        None,
        4500,
        None,
        None,
        None,
        "KZT",
        178,
        "1639000R.jpg",
        "1639000R.jpg",
        "красный",
        None,
        None,
        "да",
        "нет",
    ),
)


def _style_header(worksheet: Any, columns: Sequence[tuple[str, int]]) -> None:
    for index, (title, width) in enumerate(columns, start=1):
        cell = worksheet.cell(row=1, column=index)
        write_cell(cell, title)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        worksheet.column_dimensions[get_column_letter(index)].width = width
    worksheet.freeze_panes = "A2"


def _write_rows(worksheet: Any, rows: Sequence[Sequence[Any]], start: int = 2) -> None:
    for row_index, row in enumerate(rows, start=start):
        for column_index, value in enumerate(row, start=1):
            write_cell(worksheet.cell(row=row_index, column=column_index), value)


def _save(workbook: Workbook) -> bytes:
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _readme(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("README")
    sheet.column_dimensions["A"].width = 120
    for index, line in enumerate(README_LINES, start=1):
        cell = sheet.cell(row=index, column=1)
        write_cell(cell, line)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.cell(row=1, column=1).font = _HEADER_FONT


def build_template(*, with_examples: bool = False) -> bytes:
    """Книга для заполнения: лист «Товары», пример на отдельном листе."""
    workbook = Workbook()
    products = workbook.active
    assert products is not None
    products.title = "Товары"
    _style_header(products, TEMPLATE_COLUMNS)
    if with_examples:
        _write_rows(products, EXAMPLE_ROWS)
    else:
        # Пример — на отдельном листе: импорт его не читает, и
        # загруженный без правок шаблон не создаст демонстрационных товаров.
        example = workbook.create_sheet("Пример")
        _style_header(example, TEMPLATE_COLUMNS)
        _write_rows(example, EXAMPLE_ROWS)
    _readme(workbook)
    return _save(workbook)


def _photo(text: str, color: tuple[int, int, int]) -> bytes:
    image = Image.new("RGB", (800, 600), (245, 247, 250))
    draw = ImageDraw.Draw(image)
    draw.ellipse((250, 150, 550, 450), fill=color)
    draw.rectangle((0, 560, 800, 600), fill=(18, 76, 158))
    draw.text((20, 570), text, fill=(255, 255, 255))
    buffer = BytesIO()
    image.save(buffer, format="JPEG", quality=85)
    return buffer.getvalue()


def _scheme_pdf(title: str) -> bytes:
    """Минимальный PDF: одна страница A4, контур изделия и подпись."""
    content = (
        "0.2 w 0.07 0.3 0.62 RG 100 400 m 495 400 l S "
        "297 520 150 0 360 re S "
        "BT /F1 18 Tf 100 700 Td (" + title + ") Tj ET"
    ).encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    output = BytesIO()
    output.write(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(output.tell())
        output.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = output.tell()
    output.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        output.write(f"{offset:010d} 00000 n \n".encode())
    output.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return output.getvalue()


def build_example_zip() -> bytes:
    """Рабочий пример: книга с двумя товарами, три фото и две схемы."""
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(BUNDLE_XLSX_NAME, build_template(with_examples=True))
        archive.writestr(f"{PHOTOS_DIR}/504-176.jpg", _photo("504-176", (30, 30, 30)))
        archive.writestr(f"{PHOTOS_DIR}/504-176-2.jpg", _photo("504-176 / 2", (60, 60, 60)))
        archive.writestr(f"{PHOTOS_DIR}/1639000R.jpg", _photo("1639000R", (200, 40, 40)))
        archive.writestr(f"{SCHEMES_DIR}/504-176.pdf", _scheme_pdf("504-176 scheme"))
        archive.writestr(f"{SCHEMES_DIR}/1639000R.jpg", _photo("1639000R scheme", (18, 76, 158)))
    return buffer.getvalue()


# ------------------------------------------------------------------ отчёт
_ACTION_LABELS = {"create": "создан", "update": "обновлён", "unchanged": "без изменений"}
_PART_LABELS = {
    "slug": "адрес",
    "status": "статус",
    "is_active": "активность",
    "is_featured": "избранное",
    "unit": "единица",
    "min_order_qty": "мин. партия",
    "brand": "бренд",
    "external_id": "код 1С",
    "translations": "тексты",
    "categories": "категории",
    "attributes": "характеристики",
    "specs": "описательные характеристики",
    "pricing": "цены",
    "photos": "фото",
    "schemes": "схемы",
    "stock": "остаток",
}


def build_report(job: Any) -> bytes:
    """Отчёт импорта в XLSX: итог, ошибки, предупреждения, товары."""
    summary: dict[str, Any] = job.summary or {}
    bundle: dict[str, Any] = summary.get("bundle") or {}
    photos: dict[str, Any] = bundle.get("photos") or {}
    schemes: dict[str, Any] = bundle.get("schemes") or {}
    committed = job.status == "COMMITTED"

    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Итог"
    rows: list[tuple[str, Any]] = [
        ("Файл", job.filename),
        ("Статус", job.status),
        ("Код ошибки", job.failure_code),
        ("Создан", job.created_at.isoformat() if job.created_at else None),
        ("Применён", job.committed_at.isoformat() if job.committed_at else None),
        ("Строк в книге", job.rows_total),
        ("Товаров создано" if committed else "Будет создано товаров", job.products_created),
        ("Товаров обновлено" if committed else "Будет обновлено товаров", job.products_updated),
        ("Без изменений", job.products_unchanged),
        ("Фото в наборах товаров", photos.get("total", 0)),
        ("Фото добавлено" if committed else "Фото будет добавлено", photos.get("added", 0)),
        ("Фото снято с товаров", photos.get("removed", 0)),
        ("Новых файлов фото", photos.get("uploaded", 0)),
        ("Схем в наборах товаров", schemes.get("total", 0)),
        ("Схем добавлено" if committed else "Схем будет добавлено", schemes.get("added", 0)),
        ("Новых файлов схем", schemes.get("uploaded", 0)),
        ("Остатков обновлено", bundle.get("stock_updated", 0)),
        ("Файлов в архиве", bundle.get("files_in_archive", 0)),
        ("Неиспользованных файлов", bundle.get("files_unused", 0)),
        ("Ошибок", job.error_count),
        ("Предупреждений", job.warning_count),
    ]
    _style_header(sheet, (("Показатель", 32), ("Значение", 40)))
    _write_rows(sheet, [row for row in rows if row[1] is not None])

    issue_columns = (
        ("Лист", 14),
        ("Строка", 8),
        ("Колонка", 18),
        ("Код", 24),
        ("Сообщение", 90),
        ("Значение", 30),
    )
    for title, issues in (("Ошибки", job.errors or []), ("Предупреждения", job.warnings or [])):
        target = workbook.create_sheet(title)
        _style_header(target, issue_columns)
        _write_rows(
            target,
            [
                (
                    item.get("sheet"),
                    item.get("row"),
                    item.get("column"),
                    item.get("code"),
                    item.get("message"),
                    item.get("value"),
                )
                for item in issues
            ],
        )

    products = workbook.create_sheet("Товары")
    _style_header(
        products, (("Строка", 8), ("Артикул", 18), ("Действие", 16), ("Что меняется", 70))
    )
    _write_rows(
        products,
        [
            (
                item.get("row"),
                item.get("sku"),
                _ACTION_LABELS.get(str(item.get("action")), item.get("action")),
                ", ".join(_PART_LABELS.get(part, part) for part in item.get("changed") or []),
            )
            for item in summary.get("changes") or []
        ],
    )
    return _save(workbook)


__all__ = ["build_example_zip", "build_report", "build_template"]
