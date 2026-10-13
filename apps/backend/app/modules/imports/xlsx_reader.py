"""Безопасное чтение XLSX.

Файл пришёл от пользователя — ему не доверяется ничего:

* тип определяется по содержимому, а не по расширению: сигнатура ZIP,
  обязательные части книги и тип основной части в [Content_Types].xml.
  Старый .xls, CSV, .xlsm с макросами и зашифрованные книги отвергаются;
* XLSX — это ZIP, поэтому до разбора проверяется суммарный распакованный
  размер и число частей (защита от «ZIP-бомбы»). zipfile не отдаёт больше
  заявленного размера части, так что проверка заявленных размеров честная;
* XML разбирается через defusedxml (openpyxl подключает его сам, если
  пакет установлен): внешние сущности и «billion laughs» запрещены;
* формулы не вычисляются и не принимаются: ячейка с формулой — ошибка
  с просьбой вставить значение. Кешированному результату формулы тоже
  не доверяем: его пишет программа, создавшая файл, и он может не
  совпадать с формулой;
* на диск ничего не пишется: книга читается из памяти.

Модуль синхронный и не трогает базу — разбор выполняется в пуле
потоков, чтобы не блокировать событийный цикл.
"""

from __future__ import annotations

import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from io import BytesIO
from pathlib import PurePosixPath
from typing import Any

from openpyxl import load_workbook
from openpyxl.chartsheet import Chartsheet
from openpyxl.utils import get_column_letter

from app.core.exceptions import AppError
from app.core.logging import get_logger
from app.modules.imports.exceptions import (
    ImportFileCorruptedError,
    ImportFileTooLargeError,
    ImportUnsupportedFileError,
)
from app.modules.imports.issues import Issue, IssueCode, preview_value
from app.modules.imports.xlsx_format import COMMENT_PREFIX, XLSX_MAIN_CONTENT_TYPE, SheetSpec

logger = get_logger(__name__)

ZIP_SIGNATURE = b"PK\x03\x04"
# В обычной книге два-три десятка частей. Тысячи — признак подделки.
MAX_ZIP_ENTRIES = 1_000
MAX_CONTENT_TYPES_BYTES = 1024 * 1024
# Пустые строки (например, отформатированные до конца листа) не
# считаются строками данных, но перебирать их бесконечно тоже нельзя.
BLANK_ROWS_FACTOR = 5

CellValue = str | int | float | bool | None


@dataclass
class RowData:
    number: int
    cells: dict[str, CellValue]


@dataclass
class SheetData:
    name: str
    columns: set[str] = field(default_factory=set)
    rows: list[RowData] = field(default_factory=list)
    # Каноническое имя колонки -> заголовок, как он написан в файле.
    # Нужен отчёту: «Цена от 100», а не «price_from_100».
    labels: dict[str, str] = field(default_factory=dict)
    # Имя листа в книге («Лист1»), если оно отличается от канонического.
    title: str | None = None


@dataclass
class ParsedWorkbook:
    sheets: dict[str, SheetData]
    # Нарушения структуры (нет листа или колонки) — дальнейшая проверка
    # данных бессмысленна.
    structure_issues: list[Issue]
    # Недопустимые ячейки (формулы, ошибки Excel, даты) — строки
    # проверяются дальше, чтобы показать все ошибки файла сразу.
    cell_issues: list[Issue]


# ----------------------------------------------------------------- контейнер
def inspect_container(data: bytes, *, max_uncompressed: int) -> None:
    """Проверка ZIP-контейнера до передачи файла openpyxl."""
    if not data.startswith(ZIP_SIGNATURE):
        # Сюда попадают .xls (OLE2), CSV, PDF, картинки и зашифрованные
        # паролем книги Office — они тоже хранятся в OLE2.
        raise ImportUnsupportedFileError(details={"reason": "not_xlsx"})

    try:
        archive = zipfile.ZipFile(BytesIO(data))
    except (zipfile.BadZipFile, zipfile.LargeZipFile, ValueError, OSError) as exc:
        raise ImportFileCorruptedError(details={"reason": "bad_zip"}) from exc

    with archive:
        entries = archive.infolist()
        if len(entries) > MAX_ZIP_ENTRIES:
            raise ImportFileCorruptedError(details={"reason": "too_many_parts"})

        names: set[str] = set()
        total = 0
        for entry in entries:
            name = entry.filename
            parts = PurePosixPath(name).parts
            if name.startswith("/") or "\\" in name or ".." in parts:
                raise ImportFileCorruptedError(details={"reason": "unsafe_path"})
            if entry.flag_bits & 0x1:
                raise ImportUnsupportedFileError(details={"reason": "encrypted"})
            total += entry.file_size
            if total > max_uncompressed:
                raise ImportFileTooLargeError(
                    "Распакованное содержимое книги слишком велико",
                    details={
                        "reason": "uncompressed_size",
                        "max_uncompressed_mb": max_uncompressed // (1024 * 1024),
                    },
                )
            names.add(name)

        if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
            raise ImportUnsupportedFileError(details={"reason": "not_xlsx"})
        if any(name.lower().endswith("vbaproject.bin") for name in names):
            raise ImportUnsupportedFileError(
                "Книги с макросами не принимаются", details={"reason": "macro_enabled"}
            )

        content_types_entry = archive.getinfo("[Content_Types].xml")
        if content_types_entry.file_size > MAX_CONTENT_TYPES_BYTES:
            raise ImportFileCorruptedError(details={"reason": "bad_content_types"})
        try:
            content_types = archive.read(content_types_entry)
        except (zipfile.BadZipFile, OSError, ValueError, RuntimeError) as exc:
            raise ImportFileCorruptedError(details={"reason": "bad_zip"}) from exc
        if XLSX_MAIN_CONTENT_TYPE.encode() not in content_types:
            # .xlsm, .xltx и прочие родственники .xlsx
            raise ImportUnsupportedFileError(details={"reason": "not_xlsx"})


# ------------------------------------------------------------------ книга
def read_workbook(
    data: bytes,
    specs: dict[str, SheetSpec],
    *,
    max_rows: int,
    max_uncompressed: int,
) -> ParsedWorkbook:
    inspect_container(data, max_uncompressed=max_uncompressed)

    try:
        workbook = load_workbook(
            BytesIO(data), read_only=True, data_only=False, keep_links=False, keep_vba=False
        )
    except Exception as exc:
        # Исключения openpyxl разнообразны (KeyError, ValueError, ошибки
        # XML и defusedxml). Наружу уходит одна понятная ошибка, в лог —
        # только тип исключения, без содержимого файла.
        logger.info("import_workbook_unreadable", error=type(exc).__name__)
        raise ImportFileCorruptedError from exc

    structure: list[Issue] = []
    cells: list[Issue] = []
    sheets: dict[str, SheetData] = {}
    try:
        claimed = {spec.name for spec in specs.values()}
        for spec in specs.values():
            sheet_name = _find_sheet(workbook.sheetnames, spec, claimed)
            if sheet_name is None and spec.optional:
                continue
            if sheet_name is None:
                structure.append(
                    Issue(
                        sheet=spec.name,
                        row=None,
                        column=None,
                        code=IssueCode.MISSING_SHEET,
                        message=f"В книге нет листа «{spec.name}»",
                    )
                )
                continue
            worksheet = workbook[sheet_name]
            if isinstance(worksheet, Chartsheet):
                structure.append(
                    Issue(
                        sheet=spec.name,
                        row=None,
                        column=None,
                        code=IssueCode.NOT_A_WORKSHEET,
                        message=f"«{spec.name}» — лист диаграммы, а не таблица",
                    )
                )
                continue
            sheets[spec.name] = _read_sheet(worksheet, spec, max_rows, structure, cells)
    except AppError:
        raise
    except Exception as exc:
        logger.info("import_workbook_unreadable", error=type(exc).__name__)
        raise ImportFileCorruptedError from exc
    finally:
        workbook.close()

    return ParsedWorkbook(sheets=sheets, structure_issues=structure, cell_issues=cells)


def _find_sheet(names: list[str], spec: SheetSpec, claimed: set[str]) -> str | None:
    """Имя листа книги для спецификации.

    Точное имя — всегда. Для листа товаров пакетного импорта — ещё
    альтернативные имена и, в крайнем случае, первый лист книги, если он
    не занят другой спецификацией («Лист1» из только что созданного файла).
    """
    if spec.name in names:
        return spec.name
    lowered = {name.strip().lower(): name for name in names}
    for alternative in spec.alternatives:
        if alternative in lowered:
            return lowered[alternative]
    if spec.first_sheet_fallback and names and names[0] not in claimed:
        return names[0]
    return None


def _read_sheet(
    worksheet: Any,
    spec: SheetSpec,
    max_rows: int,
    structure: list[Issue],
    cell_issues: list[Issue],
) -> SheetData:
    # Размер листа, записанный в файле, бывает неверным (его пишет
    # программа-источник). Сбрасываем и читаем фактические строки.
    worksheet.reset_dimensions()
    rows = worksheet.iter_rows()
    sheet = SheetData(name=spec.name, title=getattr(worksheet, "title", None))

    header = next(rows, None)
    raw_headers = [cell.value for cell in header] if header else []
    if spec.resolve_header is not None:
        resolver = spec.resolve_header
        headers = [resolver(str(value)) if value is not None else None for value in raw_headers]
    else:
        headers = [_header_name(value) for value in raw_headers]
    if not any(headers):
        structure.append(
            Issue(
                sheet=spec.name,
                row=1,
                column=None,
                code=IssueCode.MISSING_HEADER,
                message="Первая строка листа должна содержать заголовки колонок",
            )
        )
        return sheet

    known = set(spec.column_names)
    columns: dict[int, str] = {}
    comments: set[int] = set()
    for index, name in enumerate(headers):
        if name is None:
            continue
        if name.startswith(COMMENT_PREFIX):
            comments.add(index)
            continue
        if name not in known and not (spec.dynamic is not None and spec.dynamic(name)):
            structure.append(
                Issue(
                    sheet=spec.name,
                    row=1,
                    column=name,
                    code=IssueCode.UNKNOWN_COLUMN,
                    message=(
                        f"Неизвестная колонка «{name}». Допустимые: "
                        f"{', '.join(spec.column_names)}. "
                        f"Колонки для пометок начинайте с «{COMMENT_PREFIX}»"
                    ),
                )
            )
            comments.add(index)
            continue
        if name in columns.values():
            structure.append(
                Issue(
                    sheet=spec.name,
                    row=1,
                    column=name,
                    code=IssueCode.DUPLICATE_COLUMN,
                    message=f"Колонка «{name}» указана дважды",
                )
            )
            comments.add(index)
            continue
        columns[index] = name
        sheet.labels[name] = str(raw_headers[index]).strip()

    sheet.columns = set(columns.values())
    present_groups = {
        column.group for column in spec.columns if column.group and column.name in sheet.columns
    }
    for column in spec.columns:
        needed = column.required or (column.group is not None and column.group in present_groups)
        if needed and column.name not in sheet.columns:
            structure.append(
                Issue(
                    sheet=spec.name,
                    row=1,
                    column=column.name,
                    code=IssueCode.MISSING_COLUMN,
                    message=f"Нет обязательной колонки «{column.name}»",
                )
            )

    physical_limit = max_rows * BLANK_ROWS_FACTOR + 1_000
    for position, row in enumerate(rows, start=2):
        if position > physical_limit:
            structure.append(_too_many_rows(spec.name, max_rows))
            break

        number = next(
            (cell.row for cell in row if getattr(cell, "row", None) is not None), position
        )
        values: dict[str, CellValue] = {}
        for index, cell in enumerate(row):
            if index in comments:
                continue
            value = _cell_value(cell, spec.name, number, headers, index, cell_issues)
            if value is None:
                continue
            column_name = columns.get(index)
            if column_name is None:
                cell_issues.append(
                    Issue(
                        sheet=spec.name,
                        row=number,
                        column=get_column_letter(index + 1),
                        code=IssueCode.VALUE_WITHOUT_HEADER,
                        message="Значение в колонке без заголовка",
                        value=preview_value(value),
                    )
                )
                continue
            values[column_name] = value

        if not values:
            continue
        if len(sheet.rows) >= max_rows:
            structure.append(_too_many_rows(spec.name, max_rows))
            break
        sheet.rows.append(RowData(number=number, cells=values))

    return sheet


def _too_many_rows(sheet: str, max_rows: int) -> Issue:
    return Issue(
        sheet=sheet,
        row=None,
        column=None,
        code=IssueCode.TOO_MANY_ROWS,
        message=f"На листе больше {max_rows} строк — разбейте файл на несколько",
    )


def _header_name(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip().lower()
    return text or None


def _cell_value(
    cell: Any,
    sheet: str,
    row: int,
    headers: list[str | None],
    index: int,
    issues: list[Issue],
) -> CellValue:
    value = cell.value
    if value is None:
        return None

    column = headers[index] if index < len(headers) and headers[index] else None
    column = column or get_column_letter(index + 1)
    data_type = getattr(cell, "data_type", None)

    if data_type == "f":
        issues.append(
            Issue(
                sheet=sheet,
                row=row,
                column=column,
                code=IssueCode.FORMULA_NOT_ALLOWED,
                message=(
                    "Ячейка содержит формулу. Замените её значением "
                    "(«Специальная вставка» -> «Значения»)"
                ),
                value=preview_value(value),
            )
        )
        return None
    if data_type == "e":
        issues.append(
            Issue(
                sheet=sheet,
                row=row,
                column=column,
                code=IssueCode.EXCEL_ERROR_VALUE,
                message="Ячейка содержит ошибку Excel",
                value=preview_value(value),
            )
        )
        return None
    if isinstance(value, datetime | date | time | timedelta):
        issues.append(
            Issue(
                sheet=sheet,
                row=row,
                column=column,
                code=IssueCode.UNSUPPORTED_VALUE,
                message="Дата или время здесь недопустимы — ожидается текст или число",
                value=preview_value(value),
            )
        )
        return None

    if isinstance(value, str):
        return value if value.strip() else None
    if isinstance(value, bool | int | float):
        return value
    return str(value)
