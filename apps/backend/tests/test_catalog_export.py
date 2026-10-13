"""Экспорт каталога в XLSX."""

from __future__ import annotations

import re
import zipfile
from decimal import Decimal
from io import BytesIO

from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncEngine

from app.modules.imports.xlsx_format import README_SHEET, XLSX_MIME
from tests.conftest import Headers
from tests.helpers import count_queries
from tests.import_helpers import (
    EXPORT,
    PRODUCTS,
    SPECS,
    TEMPLATE,
    Sheet,
    preview,
    read_sheet,
)


def sheet_xml(data: bytes, sheet: str) -> str:
    """XML листа по его имени — чтобы проверить, что именно записано в файл."""
    workbook = load_workbook(BytesIO(data), read_only=True)
    index = workbook.sheetnames.index(sheet) + 1
    workbook.close()
    with zipfile.ZipFile(BytesIO(data)) as archive:
        return archive.read(f"xl/worksheets/sheet{index}.xml").decode()


async def test_export_is_a_real_xlsx(api: AsyncClient, admin: Headers, catalog) -> None:
    response = await api.get(EXPORT, headers=admin)

    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX_MIME
    disposition = response.headers["content-disposition"]
    assert re.fullmatch(r'attachment; filename="oinau-catalog-\d{8}-\d{6}\.xlsx"', disposition)
    assert response.headers["cache-control"] == "no-store"
    assert response.content.startswith(b"PK\x03\x04")

    workbook = load_workbook(BytesIO(response.content), read_only=True)
    assert workbook.sheetnames == [*[sheet.value for sheet in Sheet], README_SHEET]
    workbook.close()


async def test_export_filename_is_readable_cross_origin(
    api: AsyncClient, admin: Headers, catalog
) -> None:
    """Админка живёт на другом источнике и читает имя файла из заголовка.

    Без Access-Control-Expose-Headers браузер скрывает Content-Disposition
    от кода страницы, и выгрузка сохраняется под запасным именем.
    """
    response = await api.get(EXPORT, headers={**admin, "Origin": "http://localhost:3001"})

    exposed = response.headers["access-control-expose-headers"]
    assert "Content-Disposition" in exposed


async def test_export_headers_match_canonical_format(
    api: AsyncClient, admin: Headers, catalog
) -> None:
    data = (await api.get(EXPORT, headers=admin)).content
    workbook = load_workbook(BytesIO(data), read_only=True)
    for sheet in Sheet:
        header = [cell.value for cell in next(workbook[sheet.value].iter_rows(max_row=1))]
        assert header == SPECS[sheet.value].column_names
    workbook.close()


async def test_export_products(api: AsyncClient, admin: Headers, catalog) -> None:
    rows = {
        row["sku"]: row
        for row in read_sheet((await api.get(EXPORT, headers=admin)).content, "Products")
    }

    # Живые товары — в том числе черновик и отключённый; удалённого нет.
    assert sorted(rows) == ["DRAFT-1", "KAN-16-PP", "KRP-M10", "OFF-1", "SID-01"]

    rope = rows["KAN-16-PP"]
    assert rope["slug"] == "kanat-16-pp"
    assert rope["name_ru"] == "Канат полипропиленовый 16 мм"
    assert rope["short_description_ru"] == "Армированный канат для лазательных элементов"
    assert rope["description_ru"] == "Подробное описание каната."
    assert rope["status"] == "PUBLISHED"
    assert rope["is_active"] is True
    assert rope["is_featured"] is True
    assert rope["unit"] == "m"
    assert rope["min_order_qty"] == 1
    assert rope["pricing_type"] == "TIERED"
    assert rope["categories"] == "kanaty"
    assert rows["OFF-1"]["is_active"] is False
    assert rows["DRAFT-1"]["status"] == "DRAFT"
    assert rows["SID-01"]["pricing_type"] == "ON_REQUEST"


async def test_export_prices_of_public_price_list_only(
    api: AsyncClient, admin: Headers, catalog
) -> None:
    rows = read_sheet((await api.get(EXPORT, headers=admin)).content, "Prices")
    tiers = [(row["sku"], row["min_qty"], row["max_qty"], row["price"]) for row in rows]

    assert tiers == [
        ("KAN-16-PP", 1, 9, 3500),
        ("KAN-16-PP", 10, 49, 3200),
        ("KAN-16-PP", 50, 99, 2900),
        ("KAN-16-PP", 100, None, 2600),
        ("KRP-M10", 1, None, 1200),
    ]
    # Дилерская цена 1111 в публичную выгрузку не попадает.
    assert all(row["price"] != 1111 for row in rows)


async def test_export_categories_attributes_and_values(
    api: AsyncClient, admin: Headers, catalog
) -> None:
    data = (await api.get(EXPORT, headers=admin)).content

    categories = {row["slug"]: row for row in read_sheet(data, "Categories")}
    assert set(categories) == {"komplektuyushchie", "kanaty", "krepleniya", "skrytaya"}
    assert categories["kanaty"]["parent_slug"] == "komplektuyushchie"
    assert categories["komplektuyushchie"]["parent_slug"] is None
    assert categories["komplektuyushchie"]["description_ru"] == "Всё для площадок"
    assert categories["skrytaya"]["is_active"] is False
    # Родители раньше потомков.
    order = [row["slug"] for row in read_sheet(data, "Categories")]
    assert order.index("komplektuyushchie") < order.index("kanaty")

    attributes = {row["code"]: row for row in read_sheet(data, "Attributes")}
    assert attributes["diameter"]["value_type"] == "NUMBER"
    assert attributes["diameter"]["unit"] == "мм"
    assert attributes["material"]["value_type"] == "OPTION"
    assert attributes["internal_note"]["is_visible"] is False
    assert attributes["certified"]["name_ru"] == "Сертифицирован"

    options = read_sheet(data, "AttributeOptions")
    assert options == [
        {"attribute_code": "material", "slug": "pp", "value_ru": "Полипропилен", "sort_order": 0}
    ]

    values = {
        (row["sku"], row["attribute_code"]): row["value"]
        for row in read_sheet(data, "ProductAttributes")
    }
    assert values == {
        ("KAN-16-PP", "diameter"): 16,
        ("KAN-16-PP", "material"): "pp",
        ("KAN-16-PP", "certified"): True,
        ("KAN-16-PP", "internal_note"): "не показывать",
    }

    specs = read_sheet(data, "ProductSpecs")
    assert specs == [{"sku": "KAN-16-PP", "name": "Страна", "value": "Казахстан"}]


async def test_export_has_no_stock(api: AsyncClient, admin: Headers, catalog) -> None:
    data = (await api.get(EXPORT, headers=admin)).content
    workbook = load_workbook(BytesIO(data), read_only=True)
    for sheet in Sheet:
        header = [
            str(cell.value).lower() for cell in next(workbook[sheet.value].iter_rows(max_row=1))
        ]
        assert not any(
            word in name for name in header for word in ("stock", "quantity", "warehouse")
        )
        for row in workbook[sheet.value].iter_rows(min_row=2, values_only=True):
            # Остаток каната в фикстуре — 152.5; в файле его быть не должно.
            assert 152.5 not in row
    workbook.close()


async def test_export_does_not_leak_storage_keys_or_internal_ids(
    api: AsyncClient, admin: Headers, catalog
) -> None:
    data = (await api.get(EXPORT, headers=admin)).content
    with zipfile.ZipFile(BytesIO(data)) as archive:
        content = " ".join(
            archive.read(name).decode("utf-8", "ignore") for name in archive.namelist()
        )
    assert "products/" not in content
    assert "original." not in content
    assert str(catalog.products["rope"].uuid) not in content


async def test_export_is_formula_injection_safe(api: AsyncClient, admin: Headers) -> None:
    dangerous = {
        "name": '=HYPERLINK("http://evil.example","click")',
        "brand": "+7 777 000 00 00",
        "short_description": "@SUM(1+1)",
        "description": "-10% скидка",
        "seo_title": "#N/A",
    }
    created = await api.post(
        PRODUCTS,
        headers=admin,
        json={
            "sku": "INJ-1",
            "slug": "inj-1",
            "brand": dangerous["brand"],
            "translations": {
                "ru": {
                    "name": dangerous["name"],
                    "short_description": dangerous["short_description"],
                    "description": dangerous["description"],
                    "seo_title": dangerous["seo_title"],
                }
            },
        },
    )
    assert created.status_code == 201, created.text

    data = (await api.get(EXPORT, headers=admin)).content

    # В XML листа нет ни одной формулы и ни одной ячейки-ошибки.
    xml = sheet_xml(data, "Products")
    assert "<f>" not in xml and "<f " not in xml
    assert 't="e"' not in xml

    workbook = load_workbook(BytesIO(data))
    worksheet = workbook["Products"]
    header = [cell.value for cell in worksheet[1]]
    row = worksheet[2]
    cells = {name: row[index] for index, name in enumerate(header)}
    for column, key in (
        ("name_ru", "name"),
        ("brand", "brand"),
        ("short_description_ru", "short_description"),
        ("description_ru", "description"),
        ("seo_title_ru", "seo_title"),
    ):
        cell = cells[column]
        # Значение не искажено (без добавленного апострофа) и хранится как текст.
        assert cell.value == dangerous[key]
        assert cell.data_type == "s"
    # Опасный первый символ -> признак «текст» (quotePrefix): Excel не
    # превратит ячейку в формулу даже после редактирования.
    assert cells["name_ru"].quotePrefix is True
    assert cells["brand"].quotePrefix is True
    assert cells["short_description_ru"].quotePrefix is True
    assert cells["description_ru"].quotePrefix is True

    # Импорт возвращает те же строки без изменений.
    report = await preview(api, admin, data)
    assert report["valid"], report["errors"]
    assert report["summary"]["products"]["unchanged"] == 1


async def test_export_query_count_does_not_grow_with_catalog(
    api: AsyncClient, admin: Headers, catalog, test_engine: AsyncEngine
) -> None:
    with count_queries(test_engine) as small:
        assert (await api.get(EXPORT, headers=admin)).status_code == 200

    category_id = catalog.categories["ropes"].id
    for index in range(12):
        await api.post(
            PRODUCTS,
            headers=admin,
            json={
                "sku": f"PERF-{index}",
                "slug": f"perf-{index}",
                "translations": {"ru": {"name": f"Товар {index}"}},
                "categories": [{"category_id": category_id}],
            },
        )

    with count_queries(test_engine) as large:
        assert (await api.get(EXPORT, headers=admin)).status_code == 200
    assert len(large) == len(small)


async def test_template_is_an_empty_canonical_workbook(api: AsyncClient, admin: Headers) -> None:
    response = await api.get(TEMPLATE, headers=admin)
    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX_MIME
    assert 'filename="oinau-catalog-template.xlsx"' in response.headers["content-disposition"]

    for sheet in Sheet:
        assert read_sheet(response.content, sheet.value) == []

    workbook = load_workbook(BytesIO(response.content), read_only=True)
    readme = [row[0] for row in workbook[README_SHEET].iter_rows(values_only=True)]
    workbook.close()
    assert any("1С" in (line or "") for line in readme)


async def test_decimals_are_exported_exactly(api: AsyncClient, admin: Headers, catalog) -> None:
    response = await api.post(
        PRODUCTS,
        headers=admin,
        json={
            "sku": "DEC-1",
            "slug": "dec-1",
            "min_order_qty": "2.125",
            "translations": {"ru": {"name": "Дробный"}},
        },
    )
    product_id = response.json()["id"]
    await api.put(
        f"{PRODUCTS}/{product_id}/pricing",
        headers=admin,
        json={
            "pricing_type": "FIXED",
            "tiers": [{"min_qty": "1", "max_qty": None, "price": "123456789012.34"}],
        },
    )
    data = (await api.get(EXPORT, headers=admin)).content

    product = next(row for row in read_sheet(data, "Products") if row["sku"] == "DEC-1")
    price = next(row for row in read_sheet(data, "Prices") if row["sku"] == "DEC-1")
    assert Decimal(repr(product["min_order_qty"])) == Decimal("2.125")
    assert Decimal(repr(price["price"])) == Decimal("123456789012.34")


async def test_export_requires_authentication(api: AsyncClient) -> None:
    assert (await api.get(EXPORT)).status_code == 401
    assert (await api.get(TEMPLATE)).status_code == 401
