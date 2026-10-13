"""Пакетный импорт каталога: ZIP с книгой, фотографиями и схемами."""

from __future__ import annotations

from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from app.modules.catalog.models import Product, ProductTranslation
from app.modules.imports.models import ImportJob
from app.modules.inventory.models import Stock
from app.modules.media.models import MediaFile, ProductImage
from app.modules.media.service_processing import MediaProcessingService
from app.modules.pricing.models import PriceTier
from app.modules.users.enums import UserRole
from tests.bundle_helpers import (
    BUNDLES,
    DWG_BYTES,
    PDF_BYTES,
    apply_bundle,
    check_bundle,
    create_job,
    import_bundle,
    make_bundle,
    messages,
    photo,
    run_apply,
    run_check,
    symlink_entry,
    workbook_bytes,
)

PUBLIC = "/api/v1/public/products"


async def _product(session: AsyncSession, sku: str) -> Product | None:
    session.expire_all()
    return (
        (
            await session.execute(
                select(Product).where(Product.sku == sku, Product.deleted_at.is_(None))
            )
        )
        .scalars()
        .first()
    )


async def _images(session: AsyncSession, product_id: int) -> list[ProductImage]:
    session.expire_all()
    return list(
        (
            await session.execute(
                select(ProductImage)
                .where(ProductImage.product_id == product_id)
                .order_by(ProductImage.sort_order)
            )
        )
        .scalars()
        .all()
    )


async def _media_count(session: AsyncSession) -> int:
    return int((await session.execute(select(func.count(MediaFile.id)))).scalar_one())


# ======================================================================= поток
async def test_new_products_with_photos_scheme_prices_and_stock(
    api, admin, catalog, object_storage, media_queue, test_engine, committed_session
):
    price_list_id = catalog.public_price_list.id
    warehouse_id = catalog.warehouse.id
    data = make_bundle(
        [
            {
                "Артикул": "504-176",
                "Название": "Круглое сиденье без подвесов",
                "Категория": "Крепления",
                "Цена от 1": 3900,
                "Цена от 100": 3650,
                "Цена от 1000": 3250,
                "Валюта": "KZT",
                "Остаток": 8376,
                "Фото": "504-176.jpg;504-176-2.jpg",
                "Схема": "504-176.pdf",
                "хар: Цвет": "чёрный",
            }
        ],
        {
            "photos/504-176.jpg": photo((10, 10, 10)),
            "photos/504-176-2.jpg": photo((20, 20, 20)),
            "schemes/504-176.pdf": PDF_BYTES,
        },
    )
    job = await check_bundle(api, admin, object_storage, test_engine, data)
    assert job["status"] == "VALID", job["errors"]
    assert job["products_created"] == 1
    assert job["bundle"]["photos"]["total"] == 2
    assert job["bundle"]["photos"]["uploaded"] == 2
    assert job["bundle"]["schemes"]["total"] == 1
    assert job["bundle"]["stock_updated"] == 1
    # Предпросмотр ничего не меняет
    assert await _product(committed_session, "504-176") is None
    assert await _media_count(committed_session) == 0

    applied = await apply_bundle(api, admin, object_storage, test_engine, job, media_queue)
    assert applied["status"] == "COMMITTED"
    assert applied["products_created"] == 1

    product = await _product(committed_session, "504-176")
    assert product is not None
    assert product.status == "PUBLISHED"
    assert product.slug == "krugloe-sidene-bez-podvesov-504-176"
    assert product.pricing_type == "TIERED"
    assert product.specs == {"Цвет": "чёрный"}

    tiers = (
        (
            await committed_session.execute(
                select(PriceTier)
                .where(
                    PriceTier.product_id == product.id,
                    PriceTier.price_list_id == price_list_id,
                )
                .order_by(PriceTier.min_qty)
            )
        )
        .scalars()
        .all()
    )
    assert [(t.min_qty, t.max_qty, t.price) for t in tiers] == [
        (Decimal("1.000"), Decimal("99.999"), Decimal("3900.00")),
        (Decimal("100.000"), Decimal("999.999"), Decimal("3650.00")),
        (Decimal("1000.000"), None, Decimal("3250.00")),
    ]

    stock = (
        await committed_session.execute(select(Stock).where(Stock.product_id == product.id))
    ).scalar_one()
    assert stock.quantity == Decimal("8376.000")
    assert stock.warehouse_id == warehouse_id
    assert stock.sync_source == "excel"

    images = await _images(committed_session, product.id)
    assert [(image.kind, image.is_main) for image in images] == [
        ("PHOTO", True),
        ("PHOTO", False),
        ("SCHEME", False),
    ]
    media = {
        item.id: item
        for item in (await committed_session.execute(select(MediaFile))).scalars().all()
    }
    assert [media[image.media_file_id].original_filename for image in images] == [
        "504-176.jpg",
        "504-176-2.jpg",
        "504-176.pdf",
    ]
    assert {item.status for item in media.values()} == {"PROCESSING"}
    assert media[images[2].media_file_id].mime_type == "application/pdf"
    # Оригиналы — в хранилище под UUID, обработка поставлена в очередь
    for item in media.values():
        assert item.storage_key in object_storage.objects
        assert item.storage_key.startswith("products/")
    assert sorted(media_queue.enqueued) == sorted(media)
    # Архив после применения удалён
    job_row = (
        await committed_session.execute(select(ImportJob).where(ImportJob.uuid == job["uuid"]))
    ).scalar_one()
    assert job_row.storage_key not in object_storage.objects


async def test_processed_scheme_and_photos_appear_on_storefront(
    api, admin, catalog, object_storage, media_queue, test_engine, committed_session
):
    data = make_bundle(
        [
            {
                "Артикул": "SEAT-1",
                "Название": "Сиденье",
                "Категория": "krepleniya",
                "Цена": 4500,
                "Фото": "seat.jpg",
                "Схема": "seat.pdf",
            }
        ],
        {"photos/seat.jpg": photo(), "schemes/seat.pdf": PDF_BYTES},
    )
    await import_bundle(api, admin, object_storage, test_engine, data, media_queue)

    factory = async_sessionmaker(test_engine, expire_on_commit=False)
    for media_id in media_queue.enqueued:
        async with factory() as session:
            result = await MediaProcessingService(session, storage=object_storage).process(media_id)
        assert result == "READY"

    product = await _product(committed_session, "SEAT-1")
    assert product is not None
    assert product.pricing_type == "FIXED"
    response = await api.get(f"{PUBLIC}/{product.slug}")
    assert response.status_code == 200, response.text
    body = response.json()
    assert [image["kind"] for image in body["images"]] == ["PHOTO"]
    assert body["price"]["type"] == "FIXED"
    assert body["price"]["price"] == "4500.00"
    [document] = body["documents"]
    assert document["kind"] == "SCHEME"
    assert document["content_type"] == "application/pdf"
    assert document["filename"] == "seat.pdf"
    assert document["url"].endswith("/variants/file.pdf")
    key = document["url"].split("/oinau-media/", 1)[-1]
    assert object_storage.dispositions[key].startswith("inline;")


# ============================================================ обновление
async def test_existing_sku_is_updated_not_duplicated(
    api, admin, catalog, object_storage, test_engine, committed_session
):
    rope_id = catalog.products["rope"].id
    data = make_bundle(
        [{"Артикул": "KAN-16-PP", "Название": "Канат ПП 16 мм, новая редакция", "Остаток": 40}]
    )
    job = await check_bundle(api, admin, object_storage, test_engine, data)
    assert job["status"] == "VALID", job["errors"]
    assert (job["products_created"], job["products_updated"]) == (0, 1)
    [change] = job["changes"]
    assert change["sku"] == "KAN-16-PP"
    assert set(change["changed"]) == {"translations", "stock"}

    await apply_bundle(api, admin, object_storage, test_engine, job)
    count = (
        await committed_session.execute(
            select(func.count(Product.id)).where(Product.sku == "KAN-16-PP")
        )
    ).scalar_one()
    assert count == 1
    product = await _product(committed_session, "KAN-16-PP")
    assert product is not None and product.id == rope_id
    name = (
        await committed_session.execute(
            select(ProductTranslation.name).where(ProductTranslation.product_id == rope_id)
        )
    ).scalar_one()
    assert name == "Канат ПП 16 мм, новая редакция"
    # Не заданные файлом части не тронуты: цены, категории, описание, характеристики
    assert product.pricing_type == "TIERED"
    assert product.specs == {"Страна": "Казахстан"}
    tiers = (
        await committed_session.execute(
            select(func.count(PriceTier.id)).where(PriceTier.product_id == rope_id)
        )
    ).scalar_one()
    assert tiers == 5  # 4 публичные + 1 дилерская
    stock = (
        await committed_session.execute(select(Stock).where(Stock.product_id == rope_id))
    ).scalar_one()
    assert stock.quantity == Decimal("40.000")
    assert stock.sync_source == "excel"


async def test_unchanged_row_is_reported_as_unchanged(
    api, admin, catalog, object_storage, test_engine
):
    data = make_bundle([{"Артикул": "KAN-16-PP", "Название": "Канат полипропиленовый 16 мм"}])
    job = await check_bundle(api, admin, object_storage, test_engine, data)
    assert job["status"] == "VALID", job["errors"]
    assert (job["products_created"], job["products_updated"], job["products_unchanged"]) == (
        0,
        0,
        1,
    )


async def test_single_price_column_makes_fixed_price(
    api, admin, catalog, object_storage, test_engine, committed_session
):
    data = make_bundle([{"Артикул": "SID-01", "Цена": "4 500,50"}])
    await import_bundle(api, admin, object_storage, test_engine, data)
    product = await _product(committed_session, "SID-01")
    assert product is not None and product.pricing_type == "FIXED"
    tiers = (
        (
            await committed_session.execute(
                select(PriceTier).where(PriceTier.product_id == product.id)
            )
        )
        .scalars()
        .all()
    )
    assert [(t.min_qty, t.max_qty, t.price) for t in tiers] == [
        (Decimal("1.000"), None, Decimal("4500.50"))
    ]


async def test_canonical_headers_and_first_sheet_fallback(
    api, admin, catalog, object_storage, test_engine, committed_session
):
    """Лист «Лист1», канонические имена колонок, как в выгрузке каталога."""
    workbook = workbook_bytes(
        [
            {
                "sku": "NEW-1",
                "name_ru": "Новая позиция",
                "categories": "kanaty",
                "pricing_type": "ON_REQUEST",
                "status": "DRAFT",
            }
        ],
        title="Лист1",
    )
    await import_bundle(api, admin, object_storage, test_engine, make_bundle(workbook=workbook))
    product = await _product(committed_session, "NEW-1")
    assert product is not None
    assert (product.status, product.pricing_type) == ("DRAFT", "ON_REQUEST")


# ============================================================== ошибки данных
async def test_row_errors_are_reported_with_row_and_sku(
    api, admin, catalog, object_storage, test_engine, committed_session
):
    data = make_bundle(
        [
            {"Артикул": "OK-1", "Название": "Позиция", "Категория": "Канаты", "Фото": "ok.jpg"},
            {"Артикул": "504-999", "Название": "Нет файла", "Фото": "504-999.jpg"},
            {
                "Артикул": "BAD-CAT",
                "Название": "Категория",
                "Категория": "Комплектующие для чего-то",
            },
            {"Артикул": "BAD SKU", "Название": "Пробел в артикуле"},
            {"Артикул": "OK-1", "Название": "Дубль"},
            {"Артикул": "NEW-NO-NAME"},
            {"Артикул": "BAD-PRICE", "Название": "Цена", "Цена": "дорого"},
            {"Артикул": "BAD-STOCK", "Название": "Остаток", "Остаток": -5},
            {"Артикул": "BAD-CUR", "Название": "Валюта", "Цена": 10, "Валюта": "USD"},
        ],
        {"photos/ok.jpg": photo()},
    )
    job = await check_bundle(api, admin, object_storage, test_engine, data)
    assert job["status"] == "INVALID"
    assert job["can_commit"] is False
    by_row = {(item["row"], item["code"]): item for item in job["errors"]}
    missing = by_row[(3, "file_not_found")]
    assert missing["message"] == "Артикул 504-999: не найден файл photos/504-999.jpg"
    assert missing["column"] == "Фото"
    assert missing["sheet"] == "Товары"
    assert by_row[(4, "category_not_found")]["message"] == (
        "Категория «Комплектующие для чего-то» не существует"
    )
    assert (5, "invalid_value") in by_row
    assert (6, "duplicate_row") in by_row
    assert by_row[(7, "required")]["message"] == "Для нового товара нужно название"
    assert (8, "invalid_value") in by_row
    assert (9, "invalid_value") in by_row
    assert (10, "currency_mismatch") in by_row
    # Архив с ошибками в хранилище не хранится
    assert not any(key.startswith("imports/") for key in object_storage.objects)
    # Каталог не менялся
    assert await _product(committed_session, "OK-1") is None


async def test_category_by_path_and_ambiguous_name(
    api, admin, catalog, object_storage, test_engine, committed_session
):
    from app.modules.catalog.models import Category, CategoryTranslation

    root_id = catalog.categories["root"].id
    other = Category(slug="drugie-kanaty", parent_id=None, depth=0, path="")
    committed_session.add(other)
    await committed_session.flush()
    other.path = str(other.id)
    committed_session.add(CategoryTranslation(category_id=other.id, locale="ru", name="Канаты"))
    await committed_session.commit()

    data = make_bundle(
        [
            {"Артикул": "A-1", "Название": "Один", "Категория": "Канаты"},
            {"Артикул": "A-2", "Название": "Два", "Категория": "Комплектующие / Канаты"},
        ]
    )
    job = await check_bundle(api, admin, object_storage, test_engine, data)
    codes = {(item["row"], item["code"]) for item in job["errors"]}
    assert codes == {(2, "category_ambiguous")}
    del root_id


async def test_unknown_column_is_rejected(api, admin, catalog, object_storage, test_engine):
    data = make_bundle([{"Артикул": "X-1", "Название": "X", "Опечатка": "да"}])
    job = await check_bundle(api, admin, object_storage, test_engine, data)
    assert job["status"] == "INVALID"
    assert [item["code"] for item in job["errors"]] == ["unknown_column"]


async def test_formula_is_rejected(api, admin, catalog, object_storage, test_engine):
    data = make_bundle([{"Артикул": "F-1", "Название": "F", "Цена": "=1+1"}])
    job = await check_bundle(api, admin, object_storage, test_engine, data)
    assert "formula_not_allowed" in [item["code"] for item in job["errors"]]


# ================================================================== файлы
async def test_invalid_and_unsupported_files(api, admin, catalog, object_storage, test_engine):
    data = make_bundle(
        [
            {"Артикул": "F-1", "Название": "Фейк", "Фото": "fake.jpg"},
            {"Артикул": "F-2", "Название": "Гиф", "Фото": "anim.gif"},
            {"Артикул": "F-3", "Название": "Схема exe", "Схема": "tool.exe"},
            {"Артикул": "F-4", "Название": "DWG", "Схема": "plan.dwg"},
        ],
        {
            "photos/fake.jpg": b"#!/bin/sh\necho pwned\n",
            "photos/anim.gif": b"GIF89a....",
            "schemes/tool.exe": b"MZ\x90\x00",
            "schemes/plan.dwg": DWG_BYTES,
        },
    )
    job = await check_bundle(api, admin, object_storage, test_engine, data)
    by_row = {item["row"]: item for item in job["errors"]}
    assert by_row[2]["code"] == "file_content_mismatch"
    assert by_row[3]["code"] == "file_unsupported"
    assert by_row[4]["code"] == "file_unsupported"
    assert 5 not in by_row  # DWG — допустимая схема


async def test_same_file_for_several_products_is_uploaded_once(
    api, admin, catalog, object_storage, test_engine, committed_session
):
    shared = photo((1, 2, 3))
    data = make_bundle(
        [
            {"Артикул": "S-1", "Название": "Один", "Фото": "shared.jpg"},
            {"Артикул": "S-2", "Название": "Два", "Фото": "shared.jpg"},
        ],
        {"photos/shared.jpg": shared},
    )
    applied = await import_bundle(api, admin, object_storage, test_engine, data)
    assert applied["bundle"]["photos"]["uploaded"] == 1
    assert applied["bundle"]["photos"]["added"] == 2
    assert await _media_count(committed_session) == 1


async def test_reimport_does_not_duplicate_files(
    api, admin, catalog, object_storage, test_engine, committed_session
):
    files = {"photos/a.jpg": photo((5, 5, 5)), "schemes/a.pdf": PDF_BYTES}
    row = {"Артикул": "R-1", "Название": "Позиция", "Фото": "a.jpg", "Схема": "a.pdf"}
    await import_bundle(api, admin, object_storage, test_engine, make_bundle([row], files))
    assert await _media_count(committed_session) == 2

    # Тот же архив ещё раз — ничего не меняется и ничего не загружается
    job = await check_bundle(api, admin, object_storage, test_engine, make_bundle([row], files))
    assert job["status"] == "VALID", job["errors"]
    assert job["products_unchanged"] == 1
    assert job["bundle"]["photos"]["uploaded"] == 0
    await apply_bundle(api, admin, object_storage, test_engine, job)
    assert await _media_count(committed_session) == 2

    # Старые файлы можно не класть в архив — достаточно имени
    job = await check_bundle(
        api,
        admin,
        object_storage,
        test_engine,
        make_bundle([{**row, "Фото": "a.jpg;b.jpg"}], {"photos/b.jpg": photo((9, 9, 9))}),
    )
    assert job["status"] == "VALID", job["errors"]
    assert job["bundle"]["photos"] == {
        "total": 2,
        "added": 1,
        "removed": 0,
        "uploaded": 1,
        "reused": 0,
    }


async def test_file_with_same_name_but_new_content_replaces_link(
    api, admin, catalog, object_storage, test_engine, committed_session
):
    row = {"Артикул": "R-2", "Название": "Позиция", "Схема": "r2.pdf", "Фото": "r2.jpg"}
    await import_bundle(
        api,
        admin,
        object_storage,
        test_engine,
        make_bundle([row], {"schemes/r2.pdf": PDF_BYTES, "photos/r2.jpg": photo((7, 7, 7))}),
    )
    product = await _product(committed_session, "R-2")
    assert product is not None
    product_id = product.id
    [old_photo, old_scheme] = [
        image.media_file_id for image in await _images(committed_session, product_id)
    ]

    new_pdf = PDF_BYTES.replace(b"%%EOF", b"% rev 2\n%%EOF")
    applied = await import_bundle(
        api,
        admin,
        object_storage,
        test_engine,
        make_bundle([row], {"schemes/r2.pdf": new_pdf}),
    )
    assert applied["bundle"]["schemes"]["added"] == 1
    assert applied["bundle"]["schemes"]["removed"] == 1
    images = await _images(committed_session, product_id)
    assert [image.kind for image in images] == ["PHOTO", "SCHEME"]
    assert images[0].media_file_id == old_photo
    assert images[1].media_file_id != old_scheme
    # Старый файл остаётся в медиатеке, но уже не привязан
    assert await _media_count(committed_session) == 3


async def test_photo_columns_and_order_define_cover(
    api, admin, catalog, object_storage, test_engine, committed_session
):
    data = make_bundle(
        [{"Артикул": "P-1", "Название": "Фото", "Фото 2": "second.png", "Фото 1": "first.jpg"}],
        {"photos/first.jpg": photo((1, 1, 1)), "photos/second.png": png_image()},
    )
    await import_bundle(api, admin, object_storage, test_engine, data)
    product = await _product(committed_session, "P-1")
    assert product is not None
    images = await _images(committed_session, product.id)
    media = {
        item.id: item.original_filename
        for item in (await committed_session.execute(select(MediaFile))).scalars().all()
    }
    assert [(media[image.media_file_id], image.is_main) for image in images] == [
        ("first.jpg", True),
        ("second.png", False),
    ]


async def test_unused_and_foreign_files_are_warnings(
    api, admin, catalog, object_storage, test_engine
):
    data = make_bundle(
        [{"Артикул": "U-1", "Название": "U"}],
        {"photos/extra.jpg": photo(), "notes.txt": b"hello", "__MACOSX/._x": b"junk"},
    )
    job = await check_bundle(api, admin, object_storage, test_engine, data)
    assert job["status"] == "VALID"
    codes = [item["code"] for item in job["warnings"]]
    assert "file_unused" in codes
    assert "file_ignored" in codes


# ========================================================= безопасность ZIP
@pytest.mark.parametrize(
    ("entries", "reason"),
    [
        ({"../evil.jpg": b"x"}, "Недопустимый путь"),
        ({"photos/../../evil.jpg": b"x"}, "Недопустимый путь"),
        ({"/etc/passwd": b"x"}, "Недопустимый путь"),
        ({"C:/windows/evil.jpg": b"x"}, "Недопустимый путь"),
    ],
)
async def test_path_traversal_is_rejected(
    api, admin, catalog, object_storage, test_engine, entries, reason
):
    data = make_bundle([{"Артикул": "T-1", "Название": "T"}], entries)
    job = await check_bundle(api, admin, object_storage, test_engine, data)
    assert job["status"] == "INVALID"
    [error] = job["errors"]
    assert error["code"] == "bundle_rejected"
    assert reason in error["message"]


async def test_symlink_is_rejected(api, admin, catalog, object_storage, test_engine):
    data = make_bundle(
        [{"Артикул": "T-1", "Название": "T"}],
        raw_entries=[symlink_entry("photos/link.jpg", "/etc/passwd")],
    )
    job = await check_bundle(api, admin, object_storage, test_engine, data)
    assert job["status"] == "INVALID"
    assert "Символические ссылки" in job["errors"][0]["message"]


async def test_not_a_zip_and_missing_workbook(api, admin, catalog, object_storage, test_engine):
    job = await check_bundle(api, admin, object_storage, test_engine, b"plain text, not zip")
    assert messages(job) == ["Файл не является ZIP-архивом"]

    job = await check_bundle(
        api, admin, object_storage, test_engine, make_bundle(files={"photos/a.jpg": photo()})
    )
    assert "нет products.xlsx" in messages(job)[0]


async def test_zip_bomb_is_rejected(api, admin, catalog, object_storage, test_engine):
    data = make_bundle(
        [{"Артикул": "B-1", "Название": "B"}], {"photos/bomb.jpg": b"\x00" * (8 * 1024 * 1024)}
    )
    job = await check_bundle(api, admin, object_storage, test_engine, data)
    assert job["status"] == "INVALID"
    assert "сжатый" in messages(job)[0]


async def test_limits_on_size_files_and_uncompressed(
    api, admin, catalog, object_storage, test_engine, monkeypatch
):
    from app.core.config import settings

    too_big = await api.post(
        BUNDLES,
        headers=admin,
        json={"filename": "big.zip", "size_bytes": settings.import_bundle_max_bytes + 1},
    )
    assert too_big.status_code == 422

    not_zip = await api.post(
        BUNDLES, headers=admin, json={"filename": "catalog.rar", "size_bytes": 10}
    )
    assert not_zip.status_code == 422

    monkeypatch.setattr(settings, "IMPORT_BUNDLE_MAX_FILES", 2)
    data = make_bundle(
        [{"Артикул": "L-1", "Название": "L"}],
        {"photos/a.jpg": photo((1, 1, 1)), "photos/b.jpg": photo((2, 2, 2))},
    )
    job = await check_bundle(api, admin, object_storage, test_engine, data)
    assert "больше 2 файлов" in messages(job)[0]

    monkeypatch.setattr(settings, "IMPORT_BUNDLE_MAX_FILES", 5000)
    monkeypatch.setattr(settings, "UPLOAD_MAX_SIZE_MB", 0)
    data = make_bundle([{"Артикул": "L-1", "Название": "L"}])
    job = await check_bundle(api, admin, object_storage, test_engine, data)
    assert job["status"] == "INVALID"


async def test_actual_archive_size_is_checked(
    api, admin, catalog, object_storage, test_engine, monkeypatch
):
    """Presigned PUT не ограничивает размер: проверка по факту на /check."""
    from app.core.config import settings

    data = make_bundle([{"Артикул": "L-1", "Название": "L"}])
    job = await create_job(api, admin, object_storage, data)
    monkeypatch.setattr(settings, "IMPORT_BUNDLE_MAX_MB", 0)
    response = await api.post(f"{BUNDLES}/{job['uuid']}/check", headers=admin)
    assert response.status_code == 202
    assert response.json()["status"] == "INVALID"
    assert not any(key.startswith("imports/") for key in object_storage.objects)


async def test_check_before_upload_is_conflict(api, admin, catalog, object_storage):
    response = await api.post(
        BUNDLES, headers=admin, json={"filename": "catalog.zip", "size_bytes": 100}
    )
    job = response.json()["job"]
    check = await api.post(f"{BUNDLES}/{job['uuid']}/check", headers=admin)
    assert check.status_code == 409
    assert check.json()["error"]["code"] == "upload_not_found"


# ============================================================ жизненный цикл
async def test_cancel_deletes_archive_and_blocks_commit(
    api, admin, catalog, object_storage, test_engine
):
    job = await check_bundle(
        api, admin, object_storage, test_engine, make_bundle([{"Артикул": "C-1", "Название": "C"}])
    )
    assert job["status"] == "VALID" and job["can_cancel"] is True
    cancelled = await api.delete(f"{BUNDLES}/{job['uuid']}", headers=admin)
    assert cancelled.status_code == 200
    assert cancelled.json()["status"] == "CANCELLED"
    assert not any(key.startswith("imports/") for key in object_storage.objects)
    commit = await api.post(f"{BUNDLES}/{job['uuid']}/commit", headers=admin)
    assert commit.status_code == 409


async def test_commit_twice_is_conflict(api, admin, catalog, object_storage, test_engine):
    applied = await import_bundle(
        api, admin, object_storage, test_engine, make_bundle([{"Артикул": "C-2", "Название": "C"}])
    )
    again = await api.post(f"{BUNDLES}/{applied['uuid']}/commit", headers=admin)
    assert again.status_code == 409
    assert again.json()["error"]["code"] == "import_already_committed"


async def test_failure_during_apply_rolls_back_everything(
    api, admin, catalog, object_storage, media_queue, test_engine, committed_session, monkeypatch
):
    from app.modules.imports import bundle_plan

    data = make_bundle(
        [
            {"Артикул": "KAN-16-PP", "Название": "Изменённое имя", "Фото": "x.jpg", "Остаток": 1},
            {"Артикул": "NEW-ROLLBACK", "Название": "Новый", "Фото": "x.jpg"},
        ],
        {"photos/x.jpg": photo((4, 4, 4))},
    )
    job = await check_bundle(api, admin, object_storage, test_engine, data)
    assert job["status"] == "VALID", job["errors"]

    async def explode(self, product_ids):
        raise RuntimeError("storage exploded")

    monkeypatch.setattr(bundle_plan.BundleApplier, "apply", explode)
    applied = await apply_bundle(api, admin, object_storage, test_engine, job, media_queue)
    assert applied["status"] == "FAILED"
    assert applied["failure_code"] == "internal_error"

    assert await _product(committed_session, "NEW-ROLLBACK") is None
    rope = await _product(committed_session, "KAN-16-PP")
    assert rope is not None
    name = (
        await committed_session.execute(
            select(ProductTranslation.name).where(ProductTranslation.product_id == rope.id)
        )
    ).scalar_one()
    assert name == "Канат полипропиленовый 16 мм"
    assert await _media_count(committed_session) == 0
    # Загруженные до сбоя файлы удалены, архив тоже
    assert not any(key.startswith(("products/", "imports/")) for key in object_storage.objects)
    assert media_queue.enqueued == []


async def test_catalog_change_between_check_and_apply_fails_revalidation(
    api, admin, catalog, object_storage, test_engine, committed_session
):
    data = make_bundle([{"Артикул": "SLUG-1", "Название": "Позиция", "slug": "zanyatyj"}])
    job = await check_bundle(api, admin, object_storage, test_engine, data)
    assert job["status"] == "VALID", job["errors"]

    committed_session.add(
        Product(sku="OTHER", slug="zanyatyj", status="PUBLISHED", pricing_type="ON_REQUEST")
    )
    await committed_session.commit()

    applied = await apply_bundle(api, admin, object_storage, test_engine, job)
    assert applied["status"] == "FAILED"
    assert applied["failure_code"] == "import_validation_failed"
    assert [item["code"] for item in applied["errors"]] == ["slug_taken"]
    assert await _product(committed_session, "SLUG-1") is None


async def test_worker_skips_jobs_in_wrong_state(api, admin, catalog, object_storage, test_engine):
    job = await create_job(
        api, admin, object_storage, make_bundle([{"Артикул": "W-1", "Название": "W"}])
    )
    # Не запрошена проверка — воркер ничего не делает
    assert await run_check(test_engine, object_storage, job["uuid"]) == "skipped"
    assert await run_apply(test_engine, object_storage, job["uuid"]) == "skipped"


# ============================================================== права и отчёт
@pytest.mark.parametrize("role", [UserRole.MANAGER, UserRole.VIEWER])
async def test_roles_without_write_cannot_upload(api, login_as, catalog, role):
    headers = await login_as(role)
    response = await api.post(
        BUNDLES, headers=headers, json={"filename": "c.zip", "size_bytes": 10}
    )
    assert response.status_code == 403


async def test_content_manager_can_import(api, login_as, catalog, object_storage, test_engine):
    headers = await login_as(UserRole.CONTENT_MANAGER)
    await import_bundle(
        api,
        headers,
        object_storage,
        test_engine,
        make_bundle([{"Артикул": "CM-1", "Название": "Позиция"}]),
    )


async def test_report_template_and_example(api, admin, catalog, object_storage, test_engine):
    job = await check_bundle(
        api,
        admin,
        object_storage,
        test_engine,
        make_bundle([{"Артикул": "REP-1", "Название": "Отчёт", "Фото": "none.jpg"}]),
    )
    report = await api.get(f"{BUNDLES}/{job['uuid']}/report", headers=admin)
    assert report.status_code == 200
    assert "attachment" in report.headers["content-disposition"]
    book = load_workbook(BytesIO(report.content))
    assert book.sheetnames == ["Итог", "Ошибки", "Предупреждения", "Товары"]
    errors = list(book["Ошибки"].iter_rows(min_row=2, values_only=True))
    assert errors[0][3] == "file_not_found"

    template = await api.get(f"{BUNDLES}/template", headers=admin)
    assert template.status_code == 200
    sheets = load_workbook(BytesIO(template.content)).sheetnames
    assert sheets[0] == "Товары"

    example = await api.get(f"{BUNDLES}/example", headers=admin)
    assert example.status_code == 200
    job = await check_bundle(api, admin, object_storage, test_engine, example.content)
    # В демо-каталоге тестов нет категории «Комплектующие для качелей» —
    # остальное (архив, файлы, цены) проходит проверку.
    assert {item["code"] for item in job["errors"]} == {"category_not_found"}


async def test_history_lists_only_bundle_jobs(api, admin, catalog, object_storage, test_engine):
    await check_bundle(
        api, admin, object_storage, test_engine, make_bundle([{"Артикул": "H-1", "Название": "H"}])
    )
    response = await api.get(BUNDLES, headers=admin)
    assert response.status_code == 200
    body = response.json()
    assert body["total"] == 1
    assert body["items"][0]["kind"] == "catalog_bundle"
    plain = await api.get("/api/v1/admin/catalog/imports", headers=admin)
    assert plain.json()["total"] == 0


async def test_cleanup_expires_old_previews(
    api, admin, catalog, object_storage, test_engine, committed_session
):
    from datetime import UTC, datetime, timedelta

    from app.modules.imports.bundle_service import BundleImportService

    job = await check_bundle(
        api, admin, object_storage, test_engine, make_bundle([{"Артикул": "E-1", "Название": "E"}])
    )
    assert any(key.startswith("imports/") for key in object_storage.objects)
    factory = async_sessionmaker(test_engine, expire_on_commit=False)
    async with factory() as session:
        count = await BundleImportService(session, storage=object_storage).cleanup(
            datetime.now(UTC) + timedelta(days=2)
        )
    assert count == 1
    refreshed = await api.get(f"{BUNDLES}/{job['uuid']}", headers=admin)
    assert refreshed.json()["status"] == "CANCELLED"
    assert not any(key.startswith("imports/") for key in object_storage.objects)


def png_image() -> bytes:
    from tests.media_helpers import png_bytes

    return png_bytes((32, 24))


def test_windows_zip_names_in_cp866_are_decoded():
    """«Сжатая папка» Windows пишет имена в OEM-кодировке без флага UTF-8."""
    import zipfile

    from app.modules.imports.bundle_archive import _decode_name

    info = zipfile.ZipInfo("x")
    info.filename = "фото/сиденье.jpg".encode("cp866").decode("cp437")
    info.flag_bits = 0
    assert _decode_name(info) == "фото/сиденье.jpg"

    utf8 = zipfile.ZipInfo("фото.jpg")
    utf8.flag_bits = 0x800
    assert _decode_name(utf8) == "фото.jpg"


async def test_russian_folder_names_are_accepted(api, admin, catalog, object_storage, test_engine):
    data = make_bundle(
        [{"Артикул": "RU-1", "Название": "Папки по-русски", "Фото": "a.jpg", "Схема": "a.pdf"}],
        {"каталог/фото/a.jpg": photo(), "каталог/схемы/a.pdf": PDF_BYTES},
        workbook_name="каталог/products.xlsx",
    )
    job = await check_bundle(api, admin, object_storage, test_engine, data)
    assert job["status"] == "VALID", job["errors"]
    assert job["bundle"]["photos"]["total"] == 1
    assert job["bundle"]["schemes"]["total"] == 1
