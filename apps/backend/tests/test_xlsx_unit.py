"""Модульные тесты разбора ячеек, записи XLSX и проверки контейнера (без БД)."""

from __future__ import annotations

import zipfile
from decimal import Decimal
from io import BytesIO
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

from app.modules.catalog.enums import ProductStatus
from app.modules.imports.cells import (
    CellError,
    as_bool,
    as_decimal,
    as_enum,
    as_int,
    as_list,
    as_text,
)
from app.modules.imports.exceptions import (
    ImportFileCorruptedError,
    ImportUnsupportedFileError,
)
from app.modules.imports.upload import sanitize_filename
from app.modules.imports.xlsx_format import build_sheet_specs
from app.modules.imports.xlsx_reader import inspect_container, read_workbook
from app.modules.imports.xlsx_writer import write_cell, write_workbook

MB = 1024 * 1024


# ================================================================== числа
@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        (0.1, Decimal("0.10")),
        (1234.56, Decimal("1234.56")),
        (3500, Decimal("3500.00")),
        (1200.5, Decimal("1200.50")),
        ("1 200,50", Decimal("1200.50")),
        ("1\u00a0200.5", Decimal("1200.50")),
        (1.500, Decimal("1.50")),
        ("2.5000", Decimal("2.50")),
    ],
)
def test_money_from_excel_values(raw: Any, expected: Decimal) -> None:
    value = as_decimal(raw, max_digits=14, decimal_places=2)
    assert value == expected
    assert str(value) == str(expected)  # масштаб — как у NUMERIC(14,2)


def test_float_is_never_used_directly() -> None:
    # Decimal(0.1) дал бы 0.1000000000000000055…; из repr(float) — ровно 0.1.
    tenth = 0.1
    assert as_decimal(tenth, max_digits=14, decimal_places=3) == Decimal("0.100")
    assert Decimal(tenth) != Decimal("0.1")
    # Результат вычисления в Excel (0.30000000000000004) не округляется молча.
    with pytest.raises(CellError, match="знаков"):
        as_decimal(0.1 + 0.2, max_digits=14, decimal_places=3)  # 0.30000000000000004
    assert as_decimal(0.3, max_digits=14, decimal_places=3) == Decimal("0.300")


@pytest.mark.parametrize("raw", ["abc", "NaN", "Infinity", True, "1e30", 10**13, None])
def test_invalid_money(raw: Any) -> None:
    with pytest.raises(CellError):
        as_decimal(raw, max_digits=14, decimal_places=2)


def test_decimal_bounds() -> None:
    with pytest.raises(CellError, match="больше 0"):
        as_decimal(0, max_digits=14, decimal_places=3, greater_than=Decimal(0))
    with pytest.raises(CellError, match="не меньше 0"):
        as_decimal(-1, max_digits=14, decimal_places=2, at_least=Decimal(0))
    assert as_decimal(99_999_999_999.999, max_digits=14, decimal_places=3) == Decimal(
        "99999999999.999"
    )


@pytest.mark.parametrize(("raw", "expected"), [(5, 5), (5.0, 5), ("7", 7), (" 12 ", 12)])
def test_integers(raw: Any, expected: int) -> None:
    assert as_int(raw) == expected


@pytest.mark.parametrize("raw", [5.5, "пять", True, 2**40])
def test_invalid_integers(raw: Any) -> None:
    with pytest.raises(CellError):
        as_int(raw)


# ============================================================ прочие типы
@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        (True, True),
        (False, False),
        (1, True),
        (0, False),
        ("да", True),
        ("НЕТ", False),
        ("true", True),
        ("FALSE", False),
        ("yes", True),
    ],
)
def test_booleans(raw: Any, expected: bool) -> None:
    assert as_bool(raw) is expected


@pytest.mark.parametrize("raw", ["может", 2, None])
def test_invalid_booleans(raw: Any) -> None:
    with pytest.raises(CellError):
        as_bool(raw)


def test_enum_is_case_insensitive() -> None:
    assert as_enum(" published ", ProductStatus) is ProductStatus.PUBLISHED
    with pytest.raises(CellError, match="DRAFT, PUBLISHED, ARCHIVED"):
        as_enum("LIVE", ProductStatus)


def test_text_rules() -> None:
    assert as_text("  Канат  ", max_length=10) == "Канат"
    assert as_text(12, max_length=10) == "12"
    assert as_text(2.5, max_length=10) == "2.5"
    assert as_text(12.0, max_length=10) == "12"
    assert as_text("диаметр < 16 мм", max_length=100) == "диаметр < 16 мм"
    assert as_text("=1+1", max_length=10) == "=1+1"  # строка, не формула
    for bad in ("<b>жирный</b>", "a\x00b", "a\u202eb"):
        with pytest.raises(CellError):
            as_text(bad, max_length=100)
    with pytest.raises(CellError, match="Перевод строки"):
        as_text("a\nb", max_length=100)
    assert as_text("a\r\nb", max_length=100, multiline=True) == "a\nb"
    with pytest.raises(CellError, match="длиннее"):
        as_text("x" * 11, max_length=10)


def test_list_cells() -> None:
    assert as_list("kanaty; krepleniya;;  ") == ["kanaty", "krepleniya"]
    assert as_list(None) == []


# ================================================================= запись
def _cell() -> Any:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    return worksheet.cell(row=1, column=1)


@pytest.mark.parametrize(
    ("value", "quote_prefix"),
    [
        ("=cmd|' /C calc'!A0", True),
        ("+7 701 000", True),
        ("-10%", True),
        ("@SUM(A1)", True),
        ("\tтаб", True),
        ("#N/A", False),
        ("Канат 16 мм", False),
    ],
)
def test_strings_are_written_as_text(value: str, quote_prefix: bool) -> None:
    cell = _cell()
    write_cell(cell, value)
    assert cell.value == value  # данные не искажены
    assert cell.data_type == "s"
    assert bool(cell.quotePrefix) is quote_prefix


def test_illegal_xml_characters_are_removed() -> None:
    cell = _cell()
    write_cell(cell, "a\x01b")
    assert cell.value == "ab"


def test_written_workbook_round_trips_values() -> None:
    specs = build_sheet_specs(["ru"], "ru")
    spec = specs["Prices"]
    data = write_workbook(
        [(spec, [["=EVIL()", Decimal("1.500"), None, Decimal("123456789012.34")]])],
        readme=["README"],
    )
    worksheet = load_workbook(BytesIO(data))["Prices"]
    assert worksheet["A2"].value == "=EVIL()"
    assert worksheet["A2"].data_type == "s"
    assert Decimal(repr(worksheet["B2"].value)) == Decimal("1.5")
    assert Decimal(repr(worksheet["D2"].value)) == Decimal("123456789012.34")


# ============================================================== контейнер
def test_container_checks() -> None:
    with pytest.raises(ImportUnsupportedFileError):
        inspect_container(b"hello", max_uncompressed=MB)
    with pytest.raises(ImportFileCorruptedError):
        inspect_container(b"PK\x03\x04broken", max_uncompressed=MB)

    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("xl/workbook.xml", "<workbook/>")
    with pytest.raises(ImportUnsupportedFileError):
        inspect_container(buffer.getvalue(), max_uncompressed=MB)


def test_reader_reports_missing_sheets_and_formulas() -> None:
    specs = build_sheet_specs(["ru"], "ru")
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Products"
    worksheet.append(specs["Products"].column_names)
    worksheet.append(["=A1"] + [None] * (len(specs["Products"].columns) - 1))
    buffer = BytesIO()
    workbook.save(buffer)

    parsed = read_workbook(buffer.getvalue(), specs, max_rows=10, max_uncompressed=MB)
    assert {issue.code for issue in parsed.structure_issues} == {"missing_sheet"}
    assert [issue.code for issue in parsed.cell_issues] == ["formula_not_allowed"]
    assert (parsed.cell_issues[0].row, parsed.cell_issues[0].column) == (2, "sku")


def test_filename_sanitizing() -> None:
    assert sanitize_filename("C:\\Users\\a\\каталог.xlsx") == "каталог.xlsx"
    assert sanitize_filename("../../etc/passwd.xlsx") == "passwd.xlsx"
    assert sanitize_filename("кат\x07алог\u202e.xlsx") == "каталог.xlsx"
    long_name = "a" * 400 + ".xlsx"
    assert len(sanitize_filename(long_name)) == 255
    assert sanitize_filename(long_name).endswith(".xlsx")
