"""Помощники тестов Excel-импорта: сборка книг, загрузка, чтение."""

from __future__ import annotations

from collections.abc import Callable
from io import BytesIO
from typing import Any

from httpx import AsyncClient, Response
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import settings
from app.modules.imports.xlsx_format import XLSX_MIME, Sheet, build_sheet_specs

SPECS = build_sheet_specs(settings.SUPPORTED_LOCALES, settings.DEFAULT_LOCALE)

EXPORT = "/api/v1/admin/catalog/export"
TEMPLATE = "/api/v1/admin/catalog/imports/template"
IMPORTS = "/api/v1/admin/catalog/imports"
PREVIEW = f"{IMPORTS}/preview"
PRODUCTS = "/api/v1/admin/products"
CATEGORIES = "/api/v1/admin/categories"
ATTRIBUTES = "/api/v1/admin/attributes"
PUBLIC_PRODUCTS = "/api/v1/public/products"


def commit_url(job_uuid: str) -> str:
    return f"{IMPORTS}/{job_uuid}/commit"


def product_row(sku: str, **overrides: Any) -> dict[str, Any]:
    """Строка листа Products со всеми обязательными значениями."""
    row: dict[str, Any] = {
        "sku": sku,
        "slug": sku.lower().replace(".", "-").replace("/", "-").replace("_", "-"),
        "name_ru": f"Товар {sku}",
        "status": "PUBLISHED",
        "is_active": True,
        "is_featured": False,
        "categories": None,
        "unit": "pcs",
        "min_order_qty": 1,
        "pricing_type": "ON_REQUEST",
    }
    row.update(overrides)
    return row


def category_row(slug: str, **overrides: Any) -> dict[str, Any]:
    row: dict[str, Any] = {
        "slug": slug,
        "parent_slug": None,
        "name_ru": f"Категория {slug}",
        "sort_order": 0,
        "is_active": True,
    }
    row.update(overrides)
    return row


def attribute_row(code: str, value_type: str, **overrides: Any) -> dict[str, Any]:
    row: dict[str, Any] = {
        "code": code,
        "value_type": value_type,
        "name_ru": f"Характеристика {code}",
        "unit": None,
        "is_filterable": True,
        "is_visible": True,
        "is_active": True,
        "sort_order": 0,
    }
    row.update(overrides)
    return row


def build_workbook(
    sheets: dict[str, list[dict[str, Any]]] | None = None,
    *,
    headers: dict[str, list[str]] | None = None,
    omit: set[str] | None = None,
) -> bytes:
    """Книга канонического формата. `headers` переопределяет заголовки листа.

    Значения пишутся через ws.append: строка «=...» станет формулой —
    ровно так, как это сделал бы Excel.
    """
    workbook = Workbook()
    default = workbook.active
    assert default is not None
    workbook.remove(default)
    for name, spec in SPECS.items():
        if omit and name in omit:
            continue
        worksheet = workbook.create_sheet(name)
        columns = (headers or {}).get(name, spec.column_names)
        worksheet.append(columns)
        for row in (sheets or {}).get(name, []):
            worksheet.append([row.get(column) for column in columns])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def read_sheet(data: bytes, name: str) -> list[dict[str, Any]]:
    workbook = load_workbook(BytesIO(data))
    worksheet = workbook[name]
    rows = list(worksheet.iter_rows(values_only=True))
    header = [str(value) for value in rows[0]]
    return [dict(zip(header, row, strict=False)) for row in rows[1:]]


def edit_workbook(data: bytes, edit: Callable[[Any], None]) -> bytes:
    """Правка книги так, как это сделал бы пользователь в Excel."""
    workbook = load_workbook(BytesIO(data))
    edit(workbook)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def set_cell(workbook: Any, sheet: str, key: tuple[str, Any], column: str, value: Any) -> None:
    """Меняет ячейку в строке, найденной по значению ключевой колонки."""
    worksheet = workbook[sheet]
    header = [cell.value for cell in worksheet[1]]
    key_index = header.index(key[0]) + 1
    column_index = header.index(column) + 1
    for row in range(2, worksheet.max_row + 1):
        if worksheet.cell(row=row, column=key_index).value == key[1]:
            worksheet.cell(row=row, column=column_index).value = value
            return
    raise AssertionError(f"Строка {key} не найдена на листе {sheet}")


def append_row(workbook: Any, sheet: str, values: dict[str, Any]) -> None:
    worksheet = workbook[sheet]
    header = [cell.value for cell in worksheet[1]]
    worksheet.append([values.get(column) for column in header])


def delete_rows(workbook: Any, sheet: str, key: tuple[str, Any]) -> None:
    worksheet = workbook[sheet]
    header = [cell.value for cell in worksheet[1]]
    key_index = header.index(key[0]) + 1
    for row in range(worksheet.max_row, 1, -1):
        if worksheet.cell(row=row, column=key_index).value == key[1]:
            worksheet.delete_rows(row)


async def upload(
    api: AsyncClient,
    url: str,
    headers: dict[str, str],
    data: bytes,
    *,
    filename: str = "catalog.xlsx",
    content_type: str = XLSX_MIME,
) -> Response:
    return await api.post(url, headers=headers, files={"file": (filename, data, content_type)})


async def preview(api: AsyncClient, headers: dict[str, str], data: bytes) -> dict[str, Any]:
    response = await upload(api, PREVIEW, headers, data)
    assert response.status_code == 200, response.text
    return response.json()


async def preview_and_commit(
    api: AsyncClient, headers: dict[str, str], data: bytes
) -> tuple[dict[str, Any], Response]:
    report = await preview(api, headers, data)
    assert report["valid"], report["errors"]
    response = await upload(api, commit_url(report["job"]["uuid"]), headers, data)
    return report, response


def errors_by(report: dict[str, Any], code: str) -> list[dict[str, Any]]:
    return [item for item in report["errors"] if item["code"] == code]


async def count_rows(session: AsyncSession, model: Any) -> int:
    return int((await session.execute(select(func.count()).select_from(model))).scalar_one())


__all__ = ["XLSX_MIME", "Sheet"]
